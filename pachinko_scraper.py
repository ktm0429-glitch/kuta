#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
パチンコ店 機種一覧スクレイパー → Excelスプレッドシート生成
対応サイト: p-town.dmm.com / p-gabu.jp
実行要件: pip install playwright openpyxl beautifulsoup4 lxml
           python -m playwright install chromium
"""

import asyncio
import re
import sys
from collections import defaultdict
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# 店舗設定
# ─────────────────────────────────────────────
SHOPS = [
    {
        "name": "第一プラザ 西浦和店",
        "source": "ptown",
        "urls": [
            "https://p-town.dmm.com/shops/saitama/3333",
        ],
    },
    {
        "name": "ガーデン与野本町",
        "source": "ptown",
        "urls": [
            "https://p-town.dmm.com/shops/saitama/3329",
        ],
    },
    {
        "name": "ミュー西浦和",
        "source": "ptown",
        "urls": [
            "https://p-town.dmm.com/shops/saitama/3331",
        ],
    },
    {
        "name": "サンライズ",
        "source": "pgabu",
        "urls": [
            "https://p-gabu.jp/hall/detail/10762",
        ],
    },
]

# スマパチ/スマスロ判定キーワード
SMART_KEYWORDS = [
    "スマパチ", "スマートパチンコ", "SMART", "スマスロ", "スマートスロット",
]

# カテゴリ正規化マッピング（サイトによって表記が異なるため）
CATEGORY_NORMALIZE = {
    # パチンコ系
    "4円パチンコ": "4円パチンコ",
    "パチンコ(4円)": "4円パチンコ",
    "4パチ": "4円パチンコ",
    "1円パチンコ": "1円パチンコ",
    "パチンコ(1円)": "1円パチンコ",
    "1パチ": "1円パチンコ",
    "2円パチンコ": "2円パチンコ",
    "0.5円パチンコ": "0.5円パチンコ",
    # スロット系
    "20円スロット": "20円スロット",
    "スロット(20円)": "20円スロット",
    "20スロ": "20円スロット",
    "5円スロット": "5円スロット",
    "スロット(5円)": "5円スロット",
    "2円スロット": "2円スロット",
}

# カテゴリ表示順（パチンコ→スロット順）
CATEGORY_ORDER = [
    "4円パチンコ",
    "2円パチンコ",
    "1円パチンコ",
    "0.5円パチンコ",
    "20円スロット",
    "5円スロット",
    "2円スロット",
]

DEBUG_HTML_DIR = Path("debug_html")


# ─────────────────────────────────────────────
# HTML取得（Playwright使用）
# ─────────────────────────────────────────────
async def fetch_page(page, url: str) -> str:
    """ページを取得してHTMLを返す"""
    print(f"  → 取得中: {url}")
    await page.goto(url, wait_until="networkidle", timeout=45000)
    # Next.js等のSPAがデータをロードするまで待機
    await asyncio.sleep(3)
    # スクロールして遅延読み込みをトリガー
    await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
    await asyncio.sleep(2)
    await page.evaluate("window.scrollTo(0, 0)")
    await asyncio.sleep(1)
    return await page.content()


# ─────────────────────────────────────────────
# p-town.dmm.com パーサー
# ─────────────────────────────────────────────
def parse_ptown(html: str, shop_name: str) -> dict:
    """
    p-town.dmm.com のHTMLを解析して機種データを返す
    返り値: {カテゴリ名: [{"name": 機種名, "count": 台数, "is_smart": bool}]}
    """
    soup = BeautifulSoup(html, "lxml")
    result = defaultdict(list)

    # ── 戦略1: data-* 属性ベースのカテゴリ検出 ──
    # p-townはReact製のため、カテゴリセクションをid/class/data属性で検索
    # よく見られるパターン:
    #   <section id="anc-pinball"> or <div class="*machine*">

    # パチンコセクション
    pinball_section = (
        soup.find(id="anc-pinball")
        or soup.find(attrs={"data-type": "pachinko"})
        or soup.find("section", class_=re.compile(r"pinball|pachinko", re.I))
    )
    # スロットセクション
    slot_section = (
        soup.find(id="anc-slot")
        or soup.find(attrs={"data-type": "slot"})
        or soup.find("section", class_=re.compile(r"slot", re.I))
    )

    sections_to_parse = []
    if pinball_section:
        sections_to_parse.append(pinball_section)
    if slot_section:
        sections_to_parse.append(slot_section)

    # セクションが見つからない場合はページ全体を対象に
    if not sections_to_parse:
        sections_to_parse = [soup]

    for section in sections_to_parse:
        # カテゴリ見出しを探す（h2, h3, h4）
        current_category = None
        for elem in section.find_all(
            ["h2", "h3", "h4", "dt", "div", "p"],
            class_=re.compile(r"category|heading|title|tit|label", re.I)
        ):
            text = elem.get_text(strip=True)
            normalized = _normalize_category(text)
            if normalized:
                current_category = normalized

            # この見出しの後の機種リストを処理
            machine_list = _find_next_machine_list(elem)
            if machine_list and current_category:
                machines = _extract_machines_from_list(machine_list)
                if machines:
                    result[current_category].extend(machines)

    # ── 戦略2: テーブル形式の検出 ──
    if not any(result.values()):
        result = _parse_table_format(soup)

    # ── 戦略3: フラットなリスト形式の検出 ──
    if not any(result.values()):
        result = _parse_flat_list(soup)

    # デバッグ: HTMLを保存
    DEBUG_HTML_DIR.mkdir(exist_ok=True)
    safe_name = re.sub(r'[^\w]', '_', shop_name)
    (DEBUG_HTML_DIR / f"{safe_name}_ptown.html").write_text(html, encoding="utf-8")
    print(f"    [DEBUG] HTML保存: debug_html/{safe_name}_ptown.html")

    return dict(result)


def _normalize_category(text: str) -> str | None:
    """カテゴリテキストを正規化"""
    for key, normalized in CATEGORY_NORMALIZE.items():
        if key in text:
            return normalized
    # パターンマッチ（例: "4円パチンコ(○台)"）
    if re.search(r'[1-9０-９]\s*円\s*パチンコ', text):
        m = re.search(r'([0-9.]+)\s*円\s*パチンコ', text)
        if m:
            return f"{m.group(1)}円パチンコ"
    if re.search(r'[1-9０-９]\s*円\s*スロット', text):
        m = re.search(r'([0-9.]+)\s*円\s*スロット', text)
        if m:
            return f"{m.group(1)}円スロット"
    return None


def _is_smart(text: str) -> bool:
    """スマパチ/スマスロ判定"""
    return any(kw in text for kw in SMART_KEYWORDS)


def _extract_count(text: str) -> int:
    """テキストから台数を抽出"""
    m = re.search(r'(\d+)\s*台', text)
    if m:
        return int(m.group(1))
    m = re.search(r'(\d+)', text)
    if m:
        return int(m.group(1))
    return 0


def _find_next_machine_list(elem):
    """見出し要素の次にある機種リストを探す"""
    next_sib = elem.find_next_sibling()
    while next_sib:
        if next_sib.name in ["ul", "ol", "table", "div"]:
            # 機種らしき要素があるか確認
            items = next_sib.find_all(["li", "tr", "div"])
            if items:
                return next_sib
        # 次のカテゴリ見出しに達したら終了
        if next_sib.name in ["h2", "h3", "h4"]:
            break
        next_sib = next_sib.find_next_sibling()
    return None


def _extract_machines_from_list(list_elem) -> list:
    """リスト要素から機種情報を抽出"""
    machines = []
    for item in list_elem.find_all(["li", "tr"]):
        text = item.get_text(separator=" ", strip=True)
        count = _extract_count(text)
        if count == 0:
            continue

        # 機種名の抽出（台数テキストを除いた部分）
        name_text = re.sub(r'\d+台', '', text).strip()
        name_text = re.sub(r'\s+', ' ', name_text).strip()

        if name_text:
            machines.append({
                "name": name_text,
                "count": count,
                "is_smart": _is_smart(text),
            })
    return machines


def _parse_table_format(soup) -> dict:
    """テーブル形式のページを解析"""
    result = defaultdict(list)
    current_category = None

    for table in soup.find_all("table"):
        for row in table.find_all("tr"):
            cells = row.find_all(["th", "td"])
            if not cells:
                continue

            row_text = " ".join(c.get_text(strip=True) for c in cells)
            norm = _normalize_category(row_text)
            if norm:
                current_category = norm
                continue

            if current_category and len(cells) >= 2:
                name = cells[0].get_text(strip=True)
                count = _extract_count(cells[-1].get_text(strip=True))
                if name and count > 0:
                    result[current_category].append({
                        "name": name,
                        "count": count,
                        "is_smart": _is_smart(row_text),
                    })
    return dict(result)


def _parse_flat_list(soup) -> dict:
    """フラットなリスト形式のページを解析"""
    result = defaultdict(list)
    current_category = None

    for elem in soup.find_all(["h1", "h2", "h3", "h4", "h5", "li", "div", "tr"]):
        text = elem.get_text(strip=True)
        norm = _normalize_category(text)
        if norm:
            current_category = norm
            continue

        if current_category:
            count = _extract_count(text)
            if count > 0 and 2 <= len(text) <= 60:
                name = re.sub(r'\d+台', '', text).strip()
                if name:
                    result[current_category].append({
                        "name": name,
                        "count": count,
                        "is_smart": _is_smart(text),
                    })
    return dict(result)


# ─────────────────────────────────────────────
# p-gabu.jp パーサー
# ─────────────────────────────────────────────
def parse_pgabu(html: str, shop_name: str) -> dict:
    """p-gabu.jp のHTMLを解析して機種データを返す"""
    soup = BeautifulSoup(html, "lxml")
    result = defaultdict(list)

    # デバッグ用HTML保存
    DEBUG_HTML_DIR.mkdir(exist_ok=True)
    safe_name = re.sub(r'[^\w]', '_', shop_name)
    (DEBUG_HTML_DIR / f"{safe_name}_pgabu.html").write_text(html, encoding="utf-8")
    print(f"    [DEBUG] HTML保存: debug_html/{safe_name}_pgabu.html")

    # p-gabu のカテゴリは "type_name" や "genre" クラスが多い
    # 戦略1: セクション見出し + リスト
    current_category = None
    for elem in soup.find_all(True):
        text = elem.get_text(strip=True)
        if len(text) > 50:
            continue

        norm = _normalize_category(text)
        if norm:
            current_category = norm
            continue

        if current_category:
            count = _extract_count(text)
            if count > 0 and 2 <= len(text) <= 60:
                name = re.sub(r'\d+台', '', text).strip()
                if name and not _normalize_category(name):
                    result[current_category].append({
                        "name": name,
                        "count": count,
                        "is_smart": _is_smart(text),
                    })

    # 戦略2: テーブル形式
    if not any(result.values()):
        result = _parse_table_format(soup)

    return dict(result)


# ─────────────────────────────────────────────
# データ集計
# ─────────────────────────────────────────────
def summarize(shop_data: dict) -> dict:
    """
    カテゴリごとに集計（スマパチ/スマスロの機種数・台数も）
    返り値:
    {
      カテゴリ名: {
        "machines": [{"name":..., "count":..., "is_smart":...}],
        "total_count": int,
        "smart_models": int,   # スマパチ/スマスロ機種数
        "smart_count": int,    # スマパチ/スマスロ台数
      }
    }
    """
    summary = {}
    for cat, machines in shop_data.items():
        total = sum(m["count"] for m in machines)
        smart_machines = [m for m in machines if m["is_smart"]]
        summary[cat] = {
            "machines": machines,
            "total_count": total,
            "smart_models": len(smart_machines),
            "smart_count": sum(m["count"] for m in smart_machines),
        }
    return summary


def _category_type(cat: str) -> str:
    """カテゴリがパチンコかスロットか返す"""
    if "パチンコ" in cat:
        return "パチンコ"
    if "スロット" in cat:
        return "スロット"
    return "その他"


def _smart_label(cat: str) -> str:
    """カテゴリに対応するスマートラベルを返す"""
    if "パチンコ" in cat:
        return "スマートパチンコ"
    if "スロット" in cat:
        return "スマートスロット"
    return "スマート"


# ─────────────────────────────────────────────
# Excel出力
# ─────────────────────────────────────────────
# 色定義
COLOR_SHOP_HEADER = "1F3864"       # 濃紺 (店舗名)
COLOR_PACHINKO_CAT = "2E75B6"     # 青 (パチンコカテゴリ)
COLOR_SLOT_CAT = "C55A11"         # オレンジ (スロットカテゴリ)
COLOR_SMART_ROW = "E2EFDA"        # 薄緑 (スマパチ/スマスロ集計行)
COLOR_HEADER_ROW = "D6E4F7"       # 薄青 (列ヘッダー)
COLOR_TOTAL_ROW = "FFF2CC"        # 薄黄 (合計行)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
MEDIUM_BORDER = Border(
    left=Side(style="medium"),
    right=Side(style="medium"),
    top=Side(style="medium"),
    bottom=Side(style="medium"),
)


def _cell(ws, row, col, value="", bold=False, font_size=11,
          bg_color=None, font_color="000000", align="left",
          border=None, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=font_size, color=font_color)
    if bg_color:
        c.fill = PatternFill(fill_type="solid", fgColor=bg_color)
    c.alignment = Alignment(
        horizontal=align,
        vertical="center",
        wrap_text=wrap,
    )
    if border:
        c.border = border
    return c


def write_excel(all_shops_data: list, output_path: str = "pachinko_report.xlsx"):
    """
    all_shops_data: [{"name": 店舗名, "summary": {カテゴリ: {...}}}, ...]
    """
    wb = Workbook()

    # ── シート1: 全店舗サマリー ──
    ws_summary = wb.active
    ws_summary.title = "全店舗サマリー"

    _write_summary_sheet(ws_summary, all_shops_data)

    # ── シート2〜: 店舗別詳細 ──
    for shop in all_shops_data:
        safe_title = shop["name"][:31]  # Excelシート名31文字制限
        ws = wb.create_sheet(title=safe_title)
        _write_shop_sheet(ws, shop)

    wb.save(output_path)
    print(f"\n✅ Excel保存完了: {output_path}")


def _write_summary_sheet(ws, all_shops_data: list):
    """全店舗サマリーシートを書く"""
    ws.freeze_panes = "A3"

    # 全カテゴリを収集してソート
    all_cats = set()
    for s in all_shops_data:
        all_cats.update(s["summary"].keys())

    def cat_sort_key(c):
        try:
            return CATEGORY_ORDER.index(c)
        except ValueError:
            return 99

    sorted_cats = sorted(all_cats, key=cat_sort_key)

    # ヘッダー行1
    ws.merge_cells("A1:B1")
    _cell(ws, 1, 1, "店舗名", bold=True, font_size=12,
          bg_color=COLOR_SHOP_HEADER, font_color="FFFFFF", align="center")

    col = 3
    cat_col_map = {}
    for cat in sorted_cats:
        ws.merge_cells(
            start_row=1, start_column=col,
            end_row=1, end_column=col + 2
        )
        cat_label = cat
        _cell(ws, 1, col, cat_label, bold=True, font_size=11,
              bg_color=COLOR_PACHINKO_CAT if "パチンコ" in cat else COLOR_SLOT_CAT,
              font_color="FFFFFF", align="center")
        cat_col_map[cat] = col
        col += 3

    # ヘッダー行2 (台数・スマ機種・スマ台数)
    _cell(ws, 2, 1, "店舗名", bold=True, bg_color=COLOR_HEADER_ROW)
    _cell(ws, 2, 2, "メモ", bold=True, bg_color=COLOR_HEADER_ROW)
    for cat, c in cat_col_map.items():
        _cell(ws, 2, c, "総台数", bold=True, bg_color=COLOR_HEADER_ROW, align="center")
        smart_lbl = "スマパチ" if "パチンコ" in cat else "スマスロ"
        _cell(ws, 2, c + 1, f"{smart_lbl}機種", bold=True,
              bg_color=COLOR_SMART_ROW, align="center")
        _cell(ws, 2, c + 2, f"{smart_lbl}台数", bold=True,
              bg_color=COLOR_SMART_ROW, align="center")

    # データ行
    for row_idx, shop in enumerate(all_shops_data, start=3):
        _cell(ws, row_idx, 1, shop["name"], bold=True)
        _cell(ws, row_idx, 2, "")

        for cat, c in cat_col_map.items():
            info = shop["summary"].get(cat)
            if info:
                _cell(ws, row_idx, c, info["total_count"], align="center")
                _cell(ws, row_idx, c + 1, info["smart_models"],
                      align="center", bg_color=COLOR_SMART_ROW)
                _cell(ws, row_idx, c + 2, info["smart_count"],
                      align="center", bg_color=COLOR_SMART_ROW)
            else:
                for ci in range(3):
                    _cell(ws, row_idx, c + ci, "-", align="center",
                          font_color="999999")

    # 列幅調整
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 10
    for cat, c in cat_col_map.items():
        for ci in range(3):
            ws.column_dimensions[get_column_letter(c + ci)].width = 12

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18


def _write_shop_sheet(ws, shop: dict):
    """店舗別詳細シートを書く"""
    shop_name = shop["name"]
    summary = shop["summary"]

    # 店舗名ヘッダー
    ws.merge_cells("A1:D1")
    _cell(ws, 1, 1, shop_name, bold=True, font_size=14,
          bg_color=COLOR_SHOP_HEADER, font_color="FFFFFF", align="center")
    ws.row_dimensions[1].height = 24

    # 列ヘッダー
    ws.merge_cells("A2:D2")
    _cell(ws, 2, 1, "", bg_color=COLOR_HEADER_ROW)

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 12

    current_row = 3

    def cat_sort_key(c):
        try:
            return CATEGORY_ORDER.index(c)
        except ValueError:
            return 99

    sorted_cats = sorted(summary.keys(), key=cat_sort_key)

    for cat in sorted_cats:
        info = summary[cat]
        machines = info["machines"]
        is_pachinko = "パチンコ" in cat
        cat_color = COLOR_PACHINKO_CAT if is_pachinko else COLOR_SLOT_CAT
        smart_lbl = "スマートパチンコ" if is_pachinko else "スマートスロット"

        # ── カテゴリヘッダー ──
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=2
        )
        _cell(ws, current_row, 1, cat, bold=True, font_size=12,
              bg_color=cat_color, font_color="FFFFFF")
        _cell(ws, current_row, 3, "台数", bold=True,
              bg_color=cat_color, font_color="FFFFFF", align="center")
        _cell(ws, current_row, 4, "スマパチ/スロット",
              bold=True, bg_color=cat_color,
              font_color="FFFFFF", align="center", font_size=9)
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        # ── 機種一覧 ──
        for m in machines:
            _cell(ws, current_row, 1, "・")
            name_cell = _cell(ws, current_row, 2, m["name"])
            count_cell = _cell(ws, current_row, 3, m["count"], align="center")
            smart_mark = "●" if m["is_smart"] else ""
            _cell(ws, current_row, 4, smart_mark,
                  align="center",
                  font_color="00AA00" if m["is_smart"] else "FFFFFF")

            # スマパチ行を薄緑に
            if m["is_smart"]:
                for col in range(1, 5):
                    ws.cell(row=current_row, column=col).fill = \
                        PatternFill(fill_type="solid", fgColor=COLOR_SMART_ROW)

            current_row += 1

        # ── カテゴリ合計 ──
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=2
        )
        _cell(ws, current_row, 1, f"合計 {len(machines)}機種",
              bold=True, bg_color=COLOR_TOTAL_ROW)
        _cell(ws, current_row, 3, info["total_count"],
              bold=True, bg_color=COLOR_TOTAL_ROW, align="center")
        current_row += 1

        # ── スマパチ/スマスロ集計 ──
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=2
        )
        smart_summary = (
            f"  {smart_lbl}: "
            f"{info['smart_models']}機種  {info['smart_count']}台"
        )
        _cell(ws, current_row, 1, smart_summary,
              bold=True, bg_color=COLOR_SMART_ROW, font_color="006600")
        _cell(ws, current_row, 3, info["smart_count"],
              bold=True, bg_color=COLOR_SMART_ROW,
              font_color="006600", align="center")
        current_row += 1

        # ── 空白行 ──
        current_row += 1


# ─────────────────────────────────────────────
# メイン処理
# ─────────────────────────────────────────────
async def main():
    try:
        from playwright.async_api import async_playwright
    except ImportError:
        print("❌ Playwright未インストール。以下を実行してください:")
        print("   pip install playwright")
        print("   python -m playwright install chromium")
        sys.exit(1)

    all_shops_data = []

    async with async_playwright() as p:
        # ブラウザ起動（ヘッドレス: UIなし）
        browser = await p.chromium.launch(
            headless=True,
            args=[
                "--no-sandbox",
                "--disable-dev-shm-usage",
                "--disable-blink-features=AutomationControlled",
            ],
        )

        # 人間っぽいブラウザコンテキスト
        context = await browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1280, "height": 800},
            locale="ja-JP",
        )

        page = await context.new_page()

        for shop in SHOPS:
            print(f"\n📍 {shop['name']} を処理中...")
            shop_html_parts = []

            for url in shop["urls"]:
                try:
                    html = await fetch_page(page, url)
                    shop_html_parts.append(html)
                except Exception as e:
                    print(f"  ⚠️  取得失敗: {url} → {e}")

            if not shop_html_parts:
                print(f"  ❌ {shop['name']}: データ取得失敗。スキップします。")
                continue

            # HTMLを結合して解析
            combined_html = "\n".join(shop_html_parts)
            if shop["source"] == "ptown":
                shop_data = parse_ptown(combined_html, shop["name"])
            else:
                shop_data = parse_pgabu(combined_html, shop["name"])

            if not shop_data:
                print(f"  ⚠️  {shop['name']}: 機種データが見つかりませんでした。")
                print(f"      debug_html/ フォルダのHTMLを確認してください。")
            else:
                total_machines = sum(
                    len(v) for v in shop_data.values()
                )
                print(f"  ✅ {total_machines}機種取得")

            summary = summarize(shop_data)
            all_shops_data.append({
                "name": shop["name"],
                "summary": summary,
            })

        await browser.close()

    if not all_shops_data:
        print("\n❌ データが取得できませんでした。")
        sys.exit(1)

    write_excel(all_shops_data, "pachinko_report.xlsx")

    # コンソールにもサマリー表示
    print("\n" + "=" * 60)
    print("📊 集計結果サマリー")
    print("=" * 60)
    for shop in all_shops_data:
        print(f"\n【{shop['name']}】")
        for cat, info in shop["summary"].items():
            smart_lbl = _smart_label(cat)
            print(f"  {cat}: {info['total_count']}台 ({len(info['machines'])}機種)")
            if info["smart_models"] > 0:
                print(
                    f"    └ {smart_lbl}: "
                    f"{info['smart_models']}機種 / {info['smart_count']}台"
                )


if __name__ == "__main__":
    asyncio.run(main())
